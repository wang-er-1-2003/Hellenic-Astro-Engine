# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

samples = [
    ('规则库-Lilly行星释义卷.xlsx', ['LIL-SAT-02', 'LIL-MAR-05']),
    ('规则库-Lilly星座释义卷.xlsx', ['LIL-SIG-LEO-01']),
    ('规则库-Lilly相位状态卷.xlsx', ['LIL-ASP-01']),
    ('规则库-Clavis基础教义卷.xlsx', None),
    ('规则库-Tetrabiblos本命卷(BookIII).xlsx', None),
]
out = []
RB = r'D:\一个电脑软件\astrology-resources\rulebooks'
for fname, want_ids in samples:
    wb = load_workbook(RB + '\\' + fname, read_only=True, data_only=True)
    ws = wb['规则条目']
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    hmap = {str(h): i for i, h in enumerate(hdr) if h}
    out.append('== %s == 列: %s' % (fname, ', '.join(str(h) for h in hdr if h)))
    shown = 0
    for r in rows[1:]:
        rid = str(r[hmap['rule_id']]) if hmap.get('rule_id') is not None and r[hmap['rule_id']] else ''
        if want_ids and rid not in want_ids:
            continue
        if shown >= 6: break
        if not want_ids and shown >= 5: break
        out.append('  %s | planets=[%s] signs=[%s] houses=[%s] aspects=[%s]' % (
            rid,
            r[hmap['planets']] if hmap.get('planets') is not None else '∅',
            r[hmap['signs']] if hmap.get('signs') is not None else '∅',
            r[hmap['houses']] if hmap.get('houses') is not None else '∅',
            r[hmap['aspects']] if hmap.get('aspects') is not None else '∅'))
        shown += 1
    wb.close()
    out.append('')
open(r'D:\一个电脑软件\星盘引擎子件\tag_samples.txt', 'w', encoding='utf-8').write('\n'.join(out))
print('ok')
