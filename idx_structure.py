# -*- coding: utf-8 -*-
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

def dump_sheet(path, name, max_rows=6, max_cols=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    lines.append('== %s ==' % os.path.basename(path))
    for ws in wb.worksheets:
        lines.append('  [sheet] %s  dims=%s' % (ws.title, ws.max_row and '%dx%d' % (ws.max_row, ws.max_column)))
    ws = wb.worksheets[0]
    lines.append('  前 %d 行:' % max_rows)
    for i, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True)):
        vals = [str(v)[:28] if v is not None else '' for v in row[:max_cols or len(row)]]
        lines.append('    R%d: %s' % (i+1, ' | '.join(vals)))
    wb.close()
    return lines

import os
out = []
idx = r'D:\一个电脑软件\astrology-resources\rulebooks\规则库-索引总表.xlsx'
out += dump_sheet(idx, '索引总表', 10)
out.append('')
out += dump_sheet(r'D:\一个电脑软件\astrology-resources\rulebooks\规则库-Lilly行星释义卷.xlsx', 'Lilly行星释义', 8)
out.append('')
out += dump_sheet(r'D:\一个电脑软件\astrology-resources\rulebooks\规则库-推运体系.xlsx', '推运体系', 8)

open(r'D:\一个电脑软件\星盘引擎子件\idx_structure.txt', 'w', encoding='utf-8').write('\n'.join(out))
print('ok')
