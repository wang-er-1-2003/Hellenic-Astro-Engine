# -*- coding: utf-8 -*-
"""第4层数据接入2：47卷全部子表(249张) → rule_tables.json（全量 + 类型标签）
标签规则按表名关键词；引擎只消费带标签的表，其余保留备用。"""
import os, re, sys, json, datetime
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

RB = r'D:\一个电脑软件\astrology-resources\rulebooks'
OUT = r'D:\一个电脑软件\星盘引擎子件\engine\data\rule_tables.json'

def tag_of(table):
    t = table
    if '友敌' in t: return 'friends_enemies'
    if '三分主' in t: return 'triplicity'
    if t.startswith('行星要素'): return 'planetary_elements'
    if '年数' in t or '周期' in t: return 'planetary_years'
    if '界terms' in t or t == '界terms' or t.startswith('界') or '埃及界' in t or '五十界' in t or '界表' in t or '界序' in t: return 'terms'
    if '面faces' in t or t == '面faces': return 'faces'
    if '庙陷擢落' in t or '落陷' in t: return 'dignity_table'
    if '擢升' in t: return 'exaltation'
    if '十二宫' in t or ('宫位' in t and '含义' in t): return 'houses_meaning'
    if '运动参数' in t: return 'planetary_motion'
    if '容许度' in t: return 'aspect_orbs'
    if '相位度数' in t: return 'aspect_orbs'
    if '行星颜色' in t or '星座颜色' in t or '宫位颜色' in t: return 'colors'
    if '阿拉伯点' in t or '点数公式' in t or '四福点' in t or ('公式' in t and '点' in t): return 'lots_formula'
    return 'other'

def clean_cell(v):
    if v is None: return ''
    s = str(v).strip()
    return s

HEADER_WORDS = ('planet', 'gender', 'sect', 'item', 'sign', 'planet(', 'house',
                '行星', '星座', '项目', '含义', '备注', '符号', '元素', '性质', '编号', '名称',
                '顺序', '度数', '日期', '术语', '相位', '宫', '类别', '条件', '说明', '年数',
                '周期', '公式', '位置', '类型', '状态', '序号', '阶段', '分配', '次序', '主星',
                '昼', '夜', '方向', '年龄', '起点', '规则', '时限', '步骤', '算法', '平衡')

def looks_like_header(row):
    """表头特征：≥2 个非空单元格，且含表头特征词。数据行（如土星要素行）不含特征词。"""
    vals = [clean_cell(c) for c in row if clean_cell(c)]
    if len(vals) < 2:
        return False
    joined = ''.join(vals)
    return any(w in joined for w in HEADER_WORDS)

def best_header(rows, max_cols):
    """表头行 = 第一个像表头的行；都不像则取前两行里非空最多的（防御性兜底）"""
    for i, r in enumerate(rows[:3]):
        if looks_like_header(r):
            return [clean_cell(c) for c in r[:max_cols]], i
    best, bestn, bi = [], -1, 0
    for i, r in enumerate(rows[:2]):
        vals = [clean_cell(c) for c in r[:max_cols]]
        n = sum(1 for v in vals if v)
        if n > bestn:
            best, bestn, bi = vals, n, i
    return best, bi

def main():
    files = sorted(f for f in os.listdir(RB) if f.lower().endswith('.xlsx') and f != '规则库-索引总表.xlsx')
    tables = []
    stats = {}
    for f in files:
        wb = load_workbook(RB + '\\' + f, read_only=True, data_only=True)
        for s in wb.sheetnames:
            if s in ('说明', '规则条目'):
                continue
            ws = wb[s]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            max_cols = max(len(r) for r in rows)
            hdr, hdr_row = best_header(rows, max_cols)
            data_rows = rows[hdr_row + 1:]
            data = []
            for r in data_rows:
                vals = [clean_cell(c) for c in r[:max_cols]]
                if any(vals):
                    data.append(vals)
            tag = tag_of(s)
            tables.append({
                'id': '%s|%s' % (f.replace('规则库-', '').replace('.xlsx', ''), s),
                'source': f,
                'table': s,
                'tag': tag,
                'header': hdr,
                'rows': data,
            })
            stats[tag] = stats.get(tag, 0) + 1
        wb.close()
    db = {
        'meta': {
            'name': '规则库子表全量（第4层数据源2）',
            'generated': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
            'total_tables': len(tables),
            'tag_stats': stats,
            'note': 'tag=other 的表为档案/判例/速查类，保留备用未消费',
        },
        'tables': tables,
    }
    with open(OUT, 'w', encoding='utf-8') as fh:
        json.dump(db, fh, ensure_ascii=False, indent=1)
    lines = ['== 子表抽取完成 ==', '总表数: %d' % len(tables), '标签分布:']
    for k, v in sorted(stats.items(), key=lambda x: -x[1]):
        lines.append('  %-20s %d' % (k, v))
    sz = os.path.getsize(OUT) / 1024
    lines.append('输出: %s (%.0f KB)' % (OUT, sz))
    open(r'D:\一个电脑软件\星盘引擎子件\build_tables_report.txt', 'w', encoding='utf-8').write('\n'.join(lines))
    print('done', len(tables), 'tables', '%.0f' % sz, 'KB')

if __name__ == '__main__':
    main()
