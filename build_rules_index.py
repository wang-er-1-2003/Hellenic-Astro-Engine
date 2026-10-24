# -*- coding: utf-8 -*-
"""第4层数据接入：47卷规则库 → rules_index.json（标签索引 + 全量规则库）"""
import os, re, sys, json, datetime
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

RB = r'D:\一个电脑软件\astrology-resources\rulebooks'
OUT = r'D:\一个电脑软件\星盘引擎子件\engine\data\rules_index.json'

# 英文标签 → 引擎中文名（与 constants.js 对齐）
EN2ZH = {
    'Sun': '太阳', 'Moon': '月亮', 'Mercury': '水星', 'Venus': '金星', 'Mars': '火星',
    'Jupiter': '木星', 'Saturn': '土星', 'Uranus': '天王星', 'Neptune': '海王星', 'Pluto': '冥王星',
    'MeanNode': '北交点', 'TrueNode': '北交点', 'NorthNode': '北交点', 'SouthNode': '南交点',
    'Chiron': '凯龙', 'PartOfFortune': '福点', 'LotOfFortune': '福点',
}
SIGN_EN2ZH = {
    'Aries': '白羊', 'Taurus': '金牛', 'Gemini': '双子', 'Cancer': '巨蟹', 'Leo': '狮子', 'Virgo': '处女',
    'Libra': '天秤', 'Scorpio': '天蝎', 'Sagittarius': '射手', 'Capricorn': '摩羯',
    'Aquarius': '宝瓶', 'Pisces': '双鱼',
    '白羊': '白羊', '金牛': '金牛', '双子': '双子', '巨蟹': '巨蟹', '狮子': '狮子', '处女': '处女',
    '天秤': '天秤', '天蝎': '天蝎', '射手': '射手', '摩羯': '摩羯', '宝瓶': '宝瓶', '双鱼': '双鱼',
    '水瓶': '宝瓶',
}
ASPECT_EN2ZH = {
    'Conjunction': '合相', 'Opposition': '对冲', 'Square': '刑', 'Trine': '三合', 'Sextile': '六合',
    '合': '合相', '冲': '对冲', '刑': '刑', '拱': '三合', '六合': '六合', '半合': '半合', '半刑': '半刑',
    '合相': '合相', '对冲': '对冲', '三分': '三合', '六分': '六合',
}

def split_tags(v):
    if not v: return []
    s = str(v).strip()
    if not s: return []
    parts = re.split(r'[,，;；、\s]+', s)
    return [p.strip() for p in parts if p.strip()]

def norm_planets(tags):
    out = []
    for t in tags:
        if t in EN2ZH: out.append(EN2ZH[t])
        elif t in EN2ZH.values(): out.append(t)
        # 7政英文别名
        elif t == 'Sol': out.append('太阳')
        elif t == 'Luna': out.append('月亮')
        else: out.append(t)
    return out

def norm_signs(tags):
    out = []
    for t in tags:
        out.append(SIGN_EN2ZH.get(t, t))
    return out

def norm_houses(tags):
    out = []
    for t in tags:
        if re.fullmatch(r'\d{1,2}', t): out.append(t)
        elif t.upper() in ('ASC', 'MC', 'DC', 'IC'): out.append(t.upper())
        elif '宫' in t:
            m = re.search(r'\d{1,2}', t)
            out.append(m.group(0) if m else t)
        else: out.append(t)
    return out

def norm_aspects(tags):
    out = []
    for t in tags:
        out.append(ASPECT_EN2ZH.get(t, t))
    return out

def read_meta(wb):
    """读说明 sheet（键值两列）"""
    meta = {}
    for ws in wb.worksheets:
        if ws.title == '说明':
            for row in ws.iter_rows(values_only=True):
                if row and row[0] is not None:
                    k = str(row[0]).strip()
                    v = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
                    meta[k] = v
            break
    return meta

def main():
    files = sorted(f for f in os.listdir(RB) if f.lower().endswith('.xlsx') and f != '规则库-索引总表.xlsx')
    rules = []
    skipped = []
    per_file = {}
    for f in files:
        p = os.path.join(RB, f)
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
            meta = read_meta(wb)
            ws = wb['规则条目'] if '规则条目' in wb.sheetnames else None
            if not ws:
                skipped.append((f, '无规则条目sheet')); wb.close(); continue
            rows = list(ws.iter_rows(values_only=True))
            # 表头定位：含 'rule_id' 的行才是真表头（部分表第1行为跨列标题，如 Bonatti4宫择时/Valens 本命总论）
            hdr_idx = 0
            for i, r in enumerate(rows):
                if r and any(str(v).strip() == 'rule_id' for v in r):
                    hdr_idx = i
                    break
            hdr = rows[hdr_idx]
            hmap = {str(h): i for i, h in enumerate(hdr) if h}
            count = 0
            for r in rows[hdr_idx + 1:]:
                if not r or all(v is None or str(v).strip() == '' for v in r):
                    continue
                def g(k): 
                    i = hmap.get(k)
                    return r[i] if i is not None and r[i] is not None else None
                rid = str(g('rule_id') or '').strip()
                entry = {
                    'rule_id': rid,
                    'source': f,
                    'book': meta.get('书名', ''),
                    'author': meta.get('作者', ''),
                    'era': meta.get('时代', ''),
                    'planets': norm_planets(split_tags(g('planets'))),
                    'signs': norm_signs(split_tags(g('signs'))),
                    'houses': norm_houses(split_tags(g('houses'))),
                    'aspects': norm_aspects(split_tags(g('aspects'))),
                    'condition': str(g('condition') or '').strip(),
                    'result': str(g('result') or '').strip(),
                    'rule_type': str(g('rule_type') or '').strip(),
                    'confidence': str(g('confidence') or '').strip(),
                    'original_text': str(g('original_text') or '').strip(),
                    'translation': str(g('translation') or '').strip(),
                    'notes': str(g('notes') or '').strip(),
                }
                rules.append(entry)
                count += 1
            per_file[f] = count
            wb.close()
        except Exception as e:
            skipped.append((f, str(e)))
    db = {
        'meta': {
            'name': '规则库全量索引（第4层数据源）',
            'generated': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
            'total_rules': len(rules),
            'source_dir': RB,
            'note': '标签列(planets/signs/houses/aspects)为英文原文规范化为中文；无标签卷(格言/判例/公式/档案)仅入全文库',
        },
        'rules': rules,
    }
    with open(OUT, 'w', encoding='utf-8') as fh:
        json.dump(db, fh, ensure_ascii=False, indent=1)
    # 统计
    tag_count = sum(1 for r in rules if r['planets'] or r['signs'] or r['houses'] or r['aspects'])
    lines = ['== 生成完成 ==', '总规则数: %d' % len(rules), '带标签规则数: %d' % tag_count, '跳过/失败: %d' % len(skipped)]
    for f, c in per_file.items():
        lines.append('  %-52s %4d' % (f, c))
    for f, e in skipped:
        lines.append('  SKIP %s: %s' % (f, e))
    sz = os.path.getsize(OUT) / 1024
    lines.append('输出: %s (%.0f KB)' % (OUT, sz))
    open(r'D:\一个电脑软件\星盘引擎子件\build_rules_report.txt', 'w', encoding='utf-8').write('\n'.join(lines))
    print('done', len(rules), 'rules,', '%.0f' % sz, 'KB')

if __name__ == '__main__':
    main()
