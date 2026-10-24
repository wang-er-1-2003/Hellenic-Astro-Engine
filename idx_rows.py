# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

path = r'D:\一个电脑软件\astrology-resources\rulebooks\规则库-索引总表.xlsx'
wb = load_workbook(path, read_only=True, data_only=True)
out = []
ws = wb['索引总表']
out.append('== 索引总表 sheet（%dx%d）== 列头+前15行:' % (ws.max_row, ws.max_column))
rows = list(ws.iter_rows(max_row=16, values_only=True))
for i, row in enumerate(rows):
    vals = [str(v)[:22] if v is not None else '' for v in row]
    out.append('  R%d: %s' % (i+1, ' | '.join(vals)))
ws2 = wb['套系清单']
out.append('')
out.append('== 套系清单（%dx%d）==:' % (ws2.max_row, ws2.max_column))
for row in list(ws2.iter_rows(max_row=12, values_only=True)):
    out.append('  ' + ' | '.join(str(v)[:40] if v is not None else '' for v in row))
ws3 = wb['方法分类']
out.append('')
out.append('== 方法分类（%dx%d）==:' % (ws3.max_row, ws3.max_column))
for row in list(ws3.iter_rows(max_row=14, values_only=True)):
    out.append('  ' + ' | '.join(str(v)[:40] if v is not None else '' for v in row))
wb.close()

# 抽一卷完整规则条目看 schema 值（Lilly 行星释义）
wb2 = load_workbook(r'D:\一个电脑软件\astrology-resources\rulebooks\规则库-Lilly行星释义卷.xlsx', read_only=True, data_only=True)
wsr = wb2['规则条目']
out.append('')
out.append('== Lilly行星释义卷 规则条目（%dx%d）列头+前4行:' % (wsr.max_row, wsr.max_column))
for i, row in enumerate(list(wsr.iter_rows(max_row=5, values_only=True))):
    out.append('  R%d: %s' % (i+1, ' | '.join(str(v)[:50] if v is not None else '' for v in row)))
# 子表一例：友敌表 / 庙陷擢落
for sn in ['庙陷擢落', '友敌表', '三分主星']:
    if sn in wb2.sheetnames:
        w = wb2[sn]
        out.append('')
        out.append('== %s（%dx%d）==:' % (sn, w.max_row, w.max_column))
        for row in list(w.iter_rows(max_row=10, values_only=True)):
            out.append('  ' + ' | '.join(str(v)[:26] if v is not None else '' for v in row))
wb2.close()
open(r'D:\一个电脑软件\星盘引擎子件\idx_rows.txt', 'w', encoding='utf-8').write('\n'.join(out))
print('ok')
