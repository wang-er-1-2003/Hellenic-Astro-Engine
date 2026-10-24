# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

RB = r'D:\一个电脑软件\astrology-resources\rulebooks'
out = []
files = sorted(f for f in os.listdir(RB) if f.lower().endswith('.xlsx') and f != '规则库-索引总表.xlsx')
for f in files:
    p = os.path.join(RB, f)
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
        sheets = wb.sheetnames
        rule_ws = None
        for s in sheets:
            if '规则' in s and '条目' in s:
                rule_ws = s; break
        if rule_ws:
            ws = wb[rule_ws]
            hdr = [str(c.value)[:14] if c.value else '' for c in next(ws.iter_rows(max_row=1))]
            nrows = ws.max_row - 1
            has_tags = all(any(h.startswith(k) for h in hdr) for k in ['planets', 'signs', 'houses', 'aspects'])
            out.append('%-52s sheets=%d 规则条目=%s rows=%d 标签列=%s' % (f, len(sheets), rule_ws, nrows, 'Y' if has_tags else 'N'))
        else:
            out.append('%-52s sheets=%d 无规则条目sheet → %s' % (f, len(sheets), sheets))
        wb.close()
    except Exception as e:
        out.append('%-52s ERROR: %s' % (f, e))
open(r'D:\一个电脑软件\星盘引擎子件\scan_rulesheets.txt', 'w', encoding='utf-8').write('\n'.join(out))
print('ok')
