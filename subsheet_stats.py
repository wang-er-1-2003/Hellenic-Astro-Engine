# -*- coding: utf-8 -*-
"""统计47卷 sheet 构成：说明 / 规则条目 / 映射子表"""
import os, sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

RB = r'D:\一个电脑软件\astrology-resources\rulebooks'
files = sorted(f for f in os.listdir(RB) if f.lower().endswith('.xlsx') and f != '规则库-索引总表.xlsx')
lines = []
total_sub = 0
total_rules = 0
total_sheets = 0
for f in files:
    wb = load_workbook(RB + '\\' + f, read_only=True)
    sheets = wb.sheetnames
    subs = [s for s in sheets if s not in ('说明', '规则条目')]
    has_rules = '规则条目' in sheets
    lines.append('%-52s sheets=%2d 子表=%2d %s' % (f, len(sheets), len(subs), ('规则条目 %d' % (wb['规则条目'].max_row - 1)) if has_rules else '无规则条目'))
    total_sub += len(subs)
    total_sheets += len(sheets)
    if has_rules: total_rules += wb['规则条目'].max_row - 1
    wb.close()
lines.insert(0, '合计: 卷数=%d 总sheet=%d 映射/子表总数=%d 规则条目总数=%d' % (len(files), total_sheets, total_sub, total_rules))
open(r'D:\一个电脑软件\星盘引擎子件\subsheet_stats.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('ok')
